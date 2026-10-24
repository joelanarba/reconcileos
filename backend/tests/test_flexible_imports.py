import io
import base64
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from sqlalchemy import select

from app.models.entities import AuditEvent, Order, Payment


def upload(client, content, name="export.csv", source="orders"):
    response = client.post(f"/api/imports/preview?source_type={source}", content=content, headers={"X-Filename": name})
    assert response.status_code == 200, response.text
    return response.json()


def remap(client, batch, **changes):
    options = batch["import_info"]["options"] | {"expected_revision": batch["import_info"]["revision"]} | changes
    response = client.post(f"/api/imports/{batch['id']}/mapping", json=options)
    assert response.status_code == 200, response.text
    return response.json()


def workbook_bytes(workbook):
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_windows_csv_and_human_headers_do_not_require_a_template(client):
    text = 'Order Number;Total (GHS);Order Date;Order Status;Description\r\nSYN-FLEX-1;"GHS 1,234.50";20/10/2026 09:30;Fulfilled;Caf\u00e9 supplies\r\n'
    batch = upload(client, text.encode("cp1252"))
    assert batch["invalid_rows"] == 0
    assert batch["import_info"]["encoding"] == "windows-1252"
    assert batch["import_info"]["mapping"]["order_ref"] == "Order Number"
    row = batch["preview"]["items"][0]
    assert row["parsed"]["expected_minor"] == "123450"
    assert row["parsed"]["ordered_at"] == "2026-10-20T09:30:00+00:00"
    assert row["parsed"]["status"] == "completed"
    assert row["raw"]["Total (GHS)"] == "GHS 1,234.50"
    assert row["raw"]["Description"] == "Caf\u00e9 supplies"


def test_utf16_tab_export_and_exact_trailing_zeros_are_supported(client):
    batch = upload(client, "Invoice No\tAmount\tDate\nSYN-UTF16\t125.5000\t20/10/2026\n".encode("utf-16"), "export.tsv")
    assert batch["invalid_rows"] == 0
    assert batch["preview"]["items"][0]["parsed"]["expected_minor"] == "12550"


def test_excel_sep_directive_is_recognized(client):
    batch = upload(client, b"sep=;\r\nOrder ID;Amount;Date\r\nSYN-SEP;12,50;20/10/2026\r\n")
    assert batch["invalid_rows"] == 0
    assert batch["preview"]["items"][0]["source_row"] == 3
    assert batch["preview"]["items"][0]["parsed"]["expected_minor"] == "1250"


def test_unknown_columns_can_be_mapped_without_rewriting_the_file(client):
    batch = upload(client, b"Ticket,Collected,Booked\nSYN-MAP,125.50,20/10/2026\n")
    assert batch["invalid_rows"] == 1
    assert client.post(f"/api/imports/{batch['id']}/confirm").status_code == 409
    mapped = remap(client, batch, column_map={"order_ref": "Ticket", "amount": "Collected", "ordered_at": "Booked"})
    assert mapped["invalid_rows"] == 0
    assert mapped["preview"]["items"][0]["raw"]["Ticket"] == "SYN-MAP"
    assert client.post(f"/api/imports/{batch['id']}/confirm", json={"expected_revision": mapped["import_info"]["revision"]}).status_code == 200


def test_duplicate_header_names_require_an_explicit_amount_selection(client):
    batch = upload(client, b"Order ID,Amount,Amount,Date\nSYN-COLUMNS,10.00,125.50,20/10/2026\n")
    assert batch["import_info"]["mapping"]["amount"] is None
    mapped = remap(client, batch, column_map={"amount": "Amount [2]"})
    assert mapped["invalid_rows"] == 0
    assert mapped["preview"]["items"][0]["parsed"]["expected_minor"] == "12550"


def test_real_xlsx_with_preamble_dates_numeric_ids_and_multiple_sheets(client):
    workbook = Workbook()
    workbook.active.title = "Notes"
    workbook.active.append(["Synthetic workbook for import testing"])
    worksheet = workbook.create_sheet("Daily orders")
    worksheet.append(["Synthetic transaction export"])
    worksheet.append([])
    worksheet.append(["Order ID", "Total (GHS)", "Order Date", "Order Status"])
    worksheet.append([1234, 125.5, datetime(2026, 10, 20, 9, 30), "PAID"])
    worksheet["A4"].number_format = "000000"
    second = workbook.create_sheet("Other orders")
    second.append(["Order ID", "Amount", "Date"])
    second.append(["SYN-OTHER", 40, datetime(2026, 10, 21)])
    batch = upload(client, workbook_bytes(workbook), "orders.xlsx")
    assert batch["import_info"]["sheet"] == "Daily orders"
    assert batch["import_info"]["header_row"] == 3
    assert batch["invalid_rows"] == 0
    row = batch["preview"]["items"][0]
    assert row["source_row"] == 4
    assert row["parsed"]["order_ref"] == "001234"
    assert row["parsed"]["expected_minor"] == "12550"
    other = remap(client, batch, sheet="Other orders", column_map={}, header_row=None)
    assert other["invalid_rows"] == 0
    assert other["preview"]["items"][0]["parsed"]["order_ref"] == "SYN-OTHER"


def test_excel_formula_is_not_executed_or_invented_and_correction_is_audited(client):
    workbook = Workbook()
    workbook.active.append(["Order ID", "Amount", "Date"])
    workbook.active.append(["SYN-FORMULA", "=100+25", datetime(2026, 10, 20)])
    batch = upload(client, workbook_bytes(workbook), "formulas.xlsx")
    assert batch["invalid_rows"] == 1
    assert "formula" in batch["preview"]["items"][0]["errors"][0]["message"]
    corrected = remap(client, batch, corrections={"2": {"amount": "125.00"}}, reason="Verified amount against the synthetic invoice.")
    assert corrected["invalid_rows"] == 0
    assert corrected["preview"]["items"][0]["raw"]["Amount"] == "=100+25"
    client.post(f"/api/imports/{batch['id']}/confirm", json={"expected_revision": corrected["import_info"]["revision"]})
    with client.app.state.sessions() as session:
        record = session.scalar(select(Order))
        assert record.expected_minor == 12500
        assert record.raw["Amount"] == "=100+25"
        event = session.scalar(select(AuditEvent))
        assert event.new_state["interpretation"]["options"]["corrections"] == {"2": {"amount": "125.00"}}


def test_explicit_exclusion_preserves_source_and_imports_only_selected_rows(client):
    batch = upload(client, b"Order ID,Amount,Date\nSYN-OK,12.50,20/10/2026\nTOTAL,not-a-transaction,20/10/2026\n")
    assert batch["invalid_rows"] == 1
    reviewed = remap(client, batch, excluded_rows=[3], reason="Statement summary is not a transaction.")
    assert reviewed["valid_rows"] == 1
    assert reviewed["invalid_rows"] == 0
    assert reviewed["total_rows"] == 2
    assert reviewed["preview"]["items"][1]["excluded"] is True
    assert client.post(f"/api/imports/{batch['id']}/confirm").status_code == 200
    with client.app.state.sessions() as session:
        assert list(session.scalars(select(Order.order_ref))) == ["SYN-OK"]


def test_payment_status_is_never_guessed_but_user_can_explain_export_values(client):
    batch = upload(client, b"Transaction ID,Order ID,Amount,Date,Network,State\nSYN-TX,SYN-ORDER,12.50,20/10/2026,Acme Pay,OKAY\n", source="payments")
    assert batch["invalid_rows"] == 1
    assert batch["import_info"]["unknown_statuses"] == ["OKAY"]
    mapped = remap(client, batch, status_map={"OKAY": "successful"})
    assert mapped["invalid_rows"] == 0
    client.post(f"/api/imports/{batch['id']}/confirm")
    with client.app.state.sessions() as session:
        record = session.scalar(select(Payment))
        assert record.provider == "ACME_PAY"
        assert record.provider_status == "successful"
        assert record.raw["State"] == "OKAY"


def test_ambiguous_separators_and_real_excess_precision_require_review(client):
    batch = upload(client, b"Order ID,Amount,Date\nSYN-DECIMAL,10.001,20/10/2026\n")
    assert batch["invalid_rows"] == 1
    assert "number format" in batch["preview"]["items"][0]["errors"][0]["message"]
    specified = remap(client, batch, decimal_style="dot")
    assert specified["invalid_rows"] == 1
    assert "fractional" in specified["preview"]["items"][0]["errors"][0]["message"]
    grouped = remap(client, specified, decimal_style="comma")
    assert grouped["invalid_rows"] == 0
    assert grouped["preview"]["items"][0]["parsed"]["expected_minor"] == "1000100"


def test_stale_interpretation_and_confirmed_file_edits_are_rejected(client):
    batch = upload(client, b"Order ID,Amount,Date\nSYN-STALE,12.50,20/10/2026\n")
    mapped = remap(client, batch, decimal_style="dot")
    stale = client.post(f"/api/imports/{batch['id']}/mapping", json=batch["import_info"]["options"] | {"expected_revision": 1})
    assert stale.status_code == 409
    assert client.post(f"/api/imports/{batch['id']}/confirm", json={"expected_revision": 1}).status_code == 409
    assert client.post(f"/api/imports/{batch['id']}/confirm", json={"expected_revision": mapped["import_info"]["revision"]}).status_code == 200
    assert client.post(f"/api/imports/{batch['id']}/mapping", json=mapped["import_info"]["options"]).status_code == 409


def test_unquoted_extra_separator_is_not_silently_a_smaller_amount(client):
    batch = upload(client, b"Order ID,Amount,Date\nSYN-BROKEN,1,250.00,20/10/2026\n")
    assert batch["invalid_rows"] == 1
    assert any(error["field"] == "row" for error in batch["preview"]["items"][0]["errors"])


def test_currency_in_source_header_is_not_silently_overridden_by_workspace_default(client):
    batch = upload(client, b"Order ID,Amount (USD),Date\nSYN-USD,12.50,20/10/2026\n")
    assert batch["invalid_rows"] == 0
    assert batch["import_info"]["options"]["default_currency"] == "USD"
    assert batch["preview"]["items"][0]["parsed"]["currency"] == "USD"
    contradictory = remap(client, batch, default_currency="GHS")
    assert contradictory["invalid_rows"] == 1
    assert any(error["field"] == "currency" for error in contradictory["preview"]["items"][0]["errors"])


def test_real_legacy_xls_values_and_dates_are_read_without_conversion_by_user(client):
    encoded = (Path(__file__).parent / "fixtures" / "orders.xls.base64").read_text()
    batch = upload(client, base64.b64decode(encoded), "legacy-orders.xls")
    assert batch["invalid_rows"] == 0
    assert batch["import_info"]["format"] == "Excel .xls"
    row = batch["preview"]["items"][0]
    assert row["parsed"]["expected_minor"] == "12550"
    assert row["parsed"]["ordered_at"] == "2026-10-20T00:00:00+00:00"
    assert row["parsed"]["order_ref"] == "SYN-XLS-001"
    assert client.post(f"/api/imports/{batch['id']}/confirm").status_code == 200


def test_nonexistent_and_ambiguous_local_times_are_not_guessed(client):
    batch = upload(client, b"Order ID,Amount,Date\nSYN-DST,12.50,2026-10-25T01:30:00\n")
    changed = remap(client, batch, timezone="Europe/London")
    assert changed["invalid_rows"] == 1
    assert "ambiguous" in changed["preview"]["items"][0]["errors"][0]["message"]


def test_corrections_and_exclusions_require_a_reason_and_real_source_rows(client):
    batch = upload(client, b"Order ID,Amount,Date\nSYN-REVIEW,12.50,20/10/2026\n")
    for changes in ({"corrections": {"2": {"amount": "20"}}}, {"excluded_rows": [2]}, {"excluded_rows": [99], "reason": "Reviewed source statement."}):
        response = client.post(f"/api/imports/{batch['id']}/mapping", json=batch["import_info"]["options"] | changes)
        assert response.status_code == 400


def test_explicit_minor_unit_heading_is_scaled_exactly(client):
    batch = upload(client, b"Order ID,Amount minor,Date\nSYN-MINOR,12550,20/10/2026\n")
    assert batch["invalid_rows"] == 0
    assert batch["import_info"]["options"]["amount_unit"] == "minor"
    assert batch["preview"]["items"][0]["parsed"]["expected_minor"] == "12550"
    assert batch["preview"]["items"][0]["edit_values"]["amount"] == "12550"


def test_currency_symbols_and_unsupported_source_units_cannot_be_assumed_as_ghs(client):
    dollar = upload(client, b"Order ID,Amount ($),Date\nSYN-DOLLAR,125.50,20/10/2026\n")
    assert dollar["invalid_rows"] == 1
    selected = remap(client, dollar, default_currency="USD")
    assert selected["invalid_rows"] == 0
    assert selected["preview"]["items"][0]["parsed"]["currency"] == "USD"
    unsupported = upload(client, b"Order ID,Amount (NGN),Date\nSYN-NGN,125.50,20/10/2026\n")
    mapped = remap(client, unsupported, column_map={"amount": "Amount (NGN)"})
    assert mapped["invalid_rows"] == 1
    assert any("unsupported currency NGN" in error["message"] for error in mapped["preview"]["items"][0]["errors"])
    compact = upload(client, b"Order ID,AmountUSD,Date\nSYN-COMPACT,125.50,20/10/2026\n")
    assert compact["invalid_rows"] == 0
    assert compact["preview"]["items"][0]["parsed"]["currency"] == "USD"
