import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal
from zipfile import BadZipFile, ZipFile

from app.core.errors import AppError


SUPPORTED_EXTENSIONS = (".csv", ".tsv", ".txt", ".xlsx", ".xls", ".xlsm")


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float, Decimal)):
        return format(Decimal(str(value)), "f")
    return str(value)


def decode_csv(content):
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        encoding = "utf-16"
    else:
        encoding = "utf-8-sig"
    try:
        decoded = content.decode(encoding)
    except UnicodeDecodeError:
        try:
            decoded = content.decode("cp1252")
            encoding = "windows-1252"
        except UnicodeDecodeError as error:
            raise AppError("invalid_encoding", "This file's text could not be read. Upload the original Excel workbook instead.") from error
    if "\x00" in decoded:
        raise AppError("invalid_file", "This does not appear to be a readable text export. Upload the original Excel workbook.")
    return decoded, encoding


def bounded_rows(rows, max_rows):
    result = []
    cell_count = 0
    for source_row, values in rows:
        if source_row > max_rows + 25:
            raise AppError("too_many_rows", f"This file exceeds the limit of {max_rows:,} transaction rows.", 413)
        if len(values) > 100:
            raise AppError("too_many_columns", "Select a worksheet with at most 100 columns.", 413)
        if any(len(value) > 10000 for value in values):
            raise AppError("large_cell", "A source cell exceeds the 10,000-character reading limit.", 413)
        cell_count += len(values)
        if cell_count > 1_000_000:
            raise AppError("large_workbook", "The workbook contains too many cells. Import a transaction worksheet separately.", 413)
        result.append({"source_row": source_row, "cells": values})
    return result


def read_text(content, max_rows):
    decoded, encoding = decode_csv(content)
    lines = decoded.splitlines()
    separator = None
    skip = 0
    if lines and re.fullmatch(r"sep=[,;\t|]", lines[0].strip(), re.IGNORECASE):
        separator = lines[0].strip()[4:]
        skip = 1
        decoded = decoded.partition("\n")[2]
    try:
        if separator is None:
            try:
                separator = csv.Sniffer().sniff(decoded[:32768], delimiters=",;\t|").delimiter
            except csv.Error:
                separator = ","
        reader = csv.reader(io.StringIO(decoded, newline=""), delimiter=separator, strict=True)
        rows = bounded_rows(((index, list(values)) for index, values in enumerate(reader, 1 + skip)), max_rows)
    except csv.Error as error:
        raise AppError("invalid_csv", "The file contains broken quoting or column boundaries. Upload its Excel version or check the reported export.") from error
    return {"format": "CSV", "encoding": encoding, "delimiter": "tab" if separator == "\t" else separator,
            "tables": [{"name": "CSV", "rows": rows, "formulas": {}, "numeric_cells": []}]}


def read_xlsx(content, max_rows):
    from openpyxl import load_workbook

    try:
        with ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(entry.file_size for entry in entries) > 41943040:
                raise AppError("large_workbook", "The expanded workbook exceeds the safe reading limit.", 413)
            if any(entry.flag_bits & 1 for entry in entries):
                raise AppError("encrypted_workbook", "Password-protected workbooks cannot be read. Upload an unprotected transaction export.")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        cached_workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        tables = []
        total_cells = 0
        try:
            if len(workbook.worksheets) > 20:
                raise AppError("large_workbook", "Import a workbook with at most 20 worksheets.", 413)
            for worksheet in workbook.worksheets:
                if (worksheet.max_column or 0) > 100 or (worksheet.max_row or 0) > max_rows + 25:
                    raise AppError("large_workbook", "The worksheet's used range exceeds the safe row or column limit. Export the transaction table separately.", 413)
                rows, formulas, numeric = [], {}, []
                cached = cached_workbook[worksheet.title]
                for row_number, (cells, cached_cells) in enumerate(zip(worksheet.iter_rows(), cached.iter_rows()), 1):
                    values = []
                    for column_number, (cell, cached_cell) in enumerate(zip(cells, cached_cells), 1):
                        coordinate = f"{row_number}:{column_number}"
                        value = cell_text(cell.value)
                        if cell.data_type == "f":
                            formulas[coordinate] = {"formula": value, "cached": cell_text(cached_cell.value)}
                        elif isinstance(cell.value, (int, float, Decimal)) and not isinstance(cell.value, bool):
                            numeric.append(coordinate)
                            if re.fullmatch(r"0{2,30}", cell.number_format or "") and Decimal(str(cell.value)) == Decimal(str(cell.value)).to_integral_value():
                                value = value.partition(".")[0].zfill(len(cell.number_format))
                        values.append(value)
                    while values and not values[-1]:
                        values.pop()
                    rows.append((row_number, values))
                    total_cells += len(values)
                    if row_number > max_rows + 25 or len(values) > 100 or total_cells > 1_000_000:
                        raise AppError("large_workbook", "The workbook exceeds the supported transaction row or cell limit.", 413)
                tables.append({"name": worksheet.title, "rows": bounded_rows(rows, max_rows), "formulas": formulas, "numeric_cells": numeric})
        finally:
            workbook.close()
            cached_workbook.close()
        return {"format": "Excel", "encoding": None, "delimiter": None, "tables": tables}
    except AppError:
        raise
    except (Exception, BadZipFile) as error:
        raise AppError("invalid_workbook", "The Excel workbook could not be read. It may be damaged or password protected.") from error


def read_xls(content, max_rows):
    import xlrd

    try:
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True, ragged_rows=True)
        tables = []
        total_cells = 0
        try:
            if workbook.nsheets > 20:
                raise AppError("large_workbook", "Import a workbook with at most 20 worksheets.", 413)
            for sheet in workbook.sheets():
                if sheet.nrows > max_rows + 25 or sheet.ncols > 100:
                    raise AppError("large_workbook", "The worksheet exceeds the supported transaction row or column limit.", 413)
                rows, numeric = [], []
                for row_index in range(sheet.nrows):
                    values = []
                    for column_index, cell in enumerate(sheet.row(row_index)):
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(cell.value, workbook.datemode).isoformat()
                        elif cell.ctype == xlrd.XL_CELL_ERROR:
                            value = "#EXCEL_ERROR"
                        else:
                            value = cell_text(cell.value)
                        if cell.ctype == xlrd.XL_CELL_NUMBER:
                            numeric.append(f"{row_index + 1}:{column_index + 1}")
                        values.append(value)
                    total_cells += len(values)
                    if total_cells > 1_000_000:
                        raise AppError("large_workbook", "The workbook contains too many cells.", 413)
                    rows.append((row_index + 1, values))
                tables.append({"name": sheet.name, "rows": bounded_rows(rows, max_rows), "formulas": {}, "numeric_cells": numeric})
        finally:
            workbook.release_resources()
        return {"format": "Excel .xls", "encoding": None, "delimiter": None, "tables": tables}
    except AppError:
        raise
    except Exception as error:
        raise AppError("invalid_workbook", "The legacy Excel workbook could not be read. It may be damaged or password protected.") from error


def read_file(content, filename, max_rows):
    lowered = filename.lower()
    if lowered.endswith((".xlsx", ".xlsm")):
        return read_xlsx(content, max_rows)
    if lowered.endswith(".xls"):
        return read_xls(content, max_rows)
    if lowered.endswith((".csv", ".tsv", ".txt")):
        return read_text(content, max_rows)
    raise AppError("unsupported_file", "Upload a CSV, TSV, XLSX, or XLS transaction file.")
